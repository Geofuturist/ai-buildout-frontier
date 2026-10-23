r"""abf-dc: cells `Epoch site x channel atlas_dcp` (checks.csv), the table for V, and the stop checks of section 2.9.
SPEC_DC_world_v0_1_0 r2, sections 2.8 and 2.9. Not run on its own.

Cell rules (the script never sets `not_found` on candidates: Atlas has an address for only part of its records):
  found       the same address (strong key: equal normalised address after the country tail is dropped)      script:dc_checks
  not_found   the site's country has 0 Atlas records (country outside the Atlas National list, mc-0002)      script:dc_checks
  found / not_found   candidates ("same place + a shared token of owner/name") judged in the table of V      human:V
  not_checked everything else (including candidates while V's table is not given)                           script:dc_checks
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import pandas as pd

import dc_lib as L
import dc_place as P

CHANNEL = "atlas_dcp"
SCRIPT = "script:dc_checks"
SEARCH_METHOD = ("atlas_dcp v1: strong key = equal normalised address (country tail dropped) -> found; country with 0 Atlas records -> not_found; "
                 "candidates = Atlas records in the same place (GeoNames/Census place, or the same US county) with a shared token of owner/name/project; "
                 "candidates are judged by V, otherwise not_checked; no fuzzy matching; Atlas copies of atlas_data_updated 2026-10-08")
GENERIC = P.STOP | set("""dc dcs campus site data center centers centre llc inc ltd corp project phase building expansion north south east west""".split())


_STREET_ABBR = {"rd": "road", "st": "street", "ave": "avenue", "av": "avenue", "blvd": "boulevard", "dr": "drive", "ln": "lane", "hwy": "highway",
                "pkwy": "parkway", "ct": "court", "cir": "circle", "pl": "place", "ter": "terrace", "n": "north", "s": "south", "e": "east", "w": "west"}
_STATE_BY_ABBR = {v: k for k, v in L.US_STATES.items()}


def _plain(s: str) -> list[str]:
    t = re.sub(r"[^a-z0-9]+", " ", __import__("unicodedata").normalize("NFKD", s).encode("ascii", "ignore").decode().casefold())
    return t.split()


def addr_key(addr: str) -> str:
    """Strong key: the address with the country tail and the postcode dropped, US state abbreviations written out, and the
    street-type / direction abbreviations of a closed table (rd -> road, st -> street, n -> north ...) spelled out. Exact otherwise."""
    toks = [t.strip() for t in (addr or "").split(",") if t.strip()]
    while toks and (toks[-1].casefold() in L.COUNTRY_TOKENS or re.fullmatch(r"[0-9]{4,6}", toks[-1])):
        toks.pop()
    out = []
    for i, t in enumerate(toks):
        m = re.fullmatch(r"([A-Za-z]{2})\s*\d{5}(?:-\d{4})?", t) or re.fullmatch(r"([A-Za-z]{2})", t)
        if m and m.group(1).upper() in _STATE_BY_ABBR and i > 0:
            out.append(_STATE_BY_ABBR[m.group(1).upper()])
            continue
        w = _plain(re.sub(r"\s+\d{4,6}$", "", t))
        out.append(" ".join(_STREET_ABBR.get(x, x) if i == 0 else x for x in w))
    return " | ".join(out)


def tokens(*texts: str, drop: set[str]) -> set[str]:
    out = set()
    for t in texts:
        for w in L.norm(t).split():
            if len(w) >= 3 and w not in GENERIC and not w.isdigit() and w not in drop:
                out.add(w)
    return out


def build_checks(epoch_out: pd.DataFrame, atlas_out: pd.DataFrame, ep: pd.DataFrame, atl: pd.DataFrame, counts: pd.DataFrame, checked_at: str,
                 v_table: Path | None, v_date: str | None, report: list[str]):
    """ep / atl: the working frames of dc_build (raw fields + record_id, iso3_source, coord_precision, unit_id, _gid, _place, _county)."""
    n_by_iso = atlas_out.groupby("iso3_source").size().to_dict()
    atl = atl.copy()
    atl["akey"] = atl["address"].map(addr_key)
    byakey: dict[str, list] = {}
    for r in atl.itertuples():
        if r.address.strip():
            byakey.setdefault(r.akey, []).append(r)
    rows, review_rows, notes = [], [], []
    n = dict(strong=0, zero=0, cand=0, rest=0)
    for _, e in ep.reset_index(drop=True).iterrows():
        rid, iso = e["record_id"], e["iso3_source"]
        status, ref, cands = "not_checked", "", []
        key = addr_key(e["Address"]) if e["Address"].strip() else ""
        strong = [r for r in byakey.get(key, []) if r.iso3_source == iso] if key else []
        if strong:
            status, ref = "found", ";".join(r.record_id for r in strong)
            n["strong"] += 1
        elif n_by_iso.get(iso, 0) == 0:
            status = "not_found"
            n["zero"] += 1
        else:
            cands = _candidates(e, atl)
            n["cand" if cands else "rest"] += 1
        rows.append(dict(record_id=rid, channel=CHANNEL, status=status, ref=ref, checked_at=checked_at, checked_by=SCRIPT, search_method=SEARCH_METHOD))
        review_rows.append(dict(epoch_record_id=rid, name=e["Name"], owner=e["Owner"], country=e["Country"], address=e["Address"], script_status=status,
                                n_candidates=len(cands),
                                candidates="\n".join(f"{c.record_id} | {c.name_or_title} | {c.primary_actor} | {c.address or '(no address)'} | {c.source_urls[:200]}" for c in cands[:40]) + (f"\n... and {len(cands) - 40} more" if len(cands) > 40 else ""),
                                verdict="", matched_ids="", note=""))
    checks, review = pd.DataFrame(rows), pd.DataFrame(review_rows)
    notes.append(f"- strong key found: {n['strong']}; country with 0 Atlas records -> not_found: {n['zero']}; rows with candidates (for V): {n['cand']}; rest not_checked: {n['rest']}")
    if v_table is not None:
        checks, vnotes = apply_v_table(checks, review, v_table, v_date, set(atlas_out["record_id"]))
        notes += vnotes
    else:
        notes.append("- V's table not given: cells with candidates stay `not_checked`; the table is written to review_atlas_dcp.xlsx")
    return checks, review, notes


def _candidates(e, atl: pd.DataFrame):
    """Atlas records in the same place as the Epoch site (same city object, or for the US the same county) with a shared token."""
    drop = set(L.norm(e["place_"]).split())
    et = tokens(e["Name"], e["Owner"], e["Project"], drop=drop)
    sub = atl[atl["iso3_source"] == e["iso3_source"]]
    cands = sub.iloc[0:0]
    if e["coord_precision"] == "city" and e["gid_"]:
        cands = sub[sub["gid_"] == e["gid_"]]
        if e["iso3_source"] == "USA" and e["county_"]:
            cands = pd.concat([cands, sub[(sub["unit_id"] == e["county_"]) | (sub["county_"] == e["county_"])]]).drop_duplicates("record_id")
    return [c for c in cands.itertuples() if et & tokens(c.name_or_title, c.primary_actor, drop=drop)]


def apply_v_table(checks: pd.DataFrame, review: pd.DataFrame, path: Path, v_date: str, atlas_ids: set[str]):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    need = ["epoch_record_id", "verdict", "matched_ids"]
    for n in need:
        if n not in hdr:
            raise L.BuildStop(f"V's table has no column {n!r}")
    ix = {n: hdr.index(n) for n in need}
    vt = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid = str(row[ix["epoch_record_id"]] or "").strip()
        if rid:
            vt[rid] = (str(row[ix["verdict"]] or "").strip().lower(), str(row[ix["matched_ids"]] or "").strip())
    if set(vt) != set(review["epoch_record_id"]):
        raise L.BuildStop(f"V's table rows differ from the 94 Epoch sites: missing {sorted(set(review['epoch_record_id']) - set(vt))[:5]}, extra {sorted(set(vt) - set(review['epoch_record_id']))[:5]}")
    when = v_date                                         # date of V's verdicts (--v-table-date), not the file time
    notes, unsure, empty = [], [], []
    checks = checks.set_index("record_id")
    n_found = n_nf = 0
    for r in review.itertuples():
        verdict, matched = vt[r.epoch_record_id]
        if r.n_candidates == 0 or r.script_status != "not_checked":
            if verdict and ((verdict == "found") != (r.script_status == "found")) and r.script_status in ("found", "not_found"):
                notes.append(f"- V's verdict {verdict!r} differs from the script cell for {r.epoch_record_id} ({r.script_status}); the script cell is kept")
            continue
        if verdict == "unsure":
            unsure.append(r.epoch_record_id)
            continue
        if verdict == "":
            empty.append(r.epoch_record_id)              # no verdict = the cell stays not_checked (not a stop)
            continue
        if verdict not in ("found", "not_found"):
            raise L.BuildStop(f"V's table: verdict {verdict!r} for {r.epoch_record_id} is not found / not_found / unsure")
        cand_ids = [ln.split(" | ")[0].strip() for ln in r.candidates.split("\n") if " | " in ln]
        if verdict == "found":
            ids = [x.strip() for x in re.split(r"[;,\s]+", matched) if x.strip()] or (cand_ids if len(cand_ids) == 1 else [])
            if not ids:
                raise L.BuildStop(f"V's table: {r.epoch_record_id} is found but matched_ids is empty and it has {len(cand_ids)} candidates")
            unknown = [i for i in ids if i not in atlas_ids]
            if unknown:
                raise L.BuildStop(f"V's table: matched_ids of {r.epoch_record_id} are not in the Atlas layer: {unknown}")
            off = [i for i in ids if i not in cand_ids]
            if off:
                notes.append(f"- {r.epoch_record_id}: matched_ids {off} are in the Atlas layer but were not among the script candidates")
            checks.loc[r.epoch_record_id, ["status", "ref"]] = ["found", ";".join(ids)]
            n_found += 1
        else:
            checks.loc[r.epoch_record_id, ["status", "ref"]] = ["not_found", ""]
            n_nf += 1
        checks.loc[r.epoch_record_id, ["checked_by", "checked_at"]] = ["human:V", when]
    if unsure:
        raise L.BuildStop(f"V's table still has `unsure` for {len(unsure)} site(s): {unsure[:10]}")
    notes.append(f"- V's table applied ({path.name}, checked_at {when}): found {n_found}, not_found {n_nf}, empty verdict (stays not_checked) {len(empty)}: {', '.join(empty)}; checked_by human:V")
    return checks.reset_index(), notes


def write_review_xlsx(review: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "review_atlas_dcp"
    cols = list(review.columns)
    ws.append(cols)
    for r in review.itertuples(index=False):
        ws.append(list(r))
    widths = dict(epoch_record_id=26, name=30, owner=24, country=14, address=40, script_status=12, n_candidates=8, candidates=90, verdict=12, matched_ids=30, note=30)
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(c, 16)
        ws.cell(1, i).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    vcol = cols.index("verdict") + 1
    dv = DataValidation(type="list", formula1='"found,not_found,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    L_ = ws.cell(1, vcol).column_letter
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, cols.index("n_candidates") + 1).value and ws.cell(r, cols.index("script_status") + 1).value == "not_checked":
            dv.add(f"{L_}{r}")
            ws.cell(r, vcol).fill = yellow
    ws2 = wb.create_sheet("how_to")
    for line in ["Verdict is needed only for yellow rows (the script found candidates: same place + a shared token).",
                 "found = the Epoch site is one of the candidate Atlas records; put its record_id(s) in matched_ids (separated by ;). If there is one candidate, matched_ids may stay empty.",
                 "not_found = none of the candidates is the site. unsure must be closed before the build.",
                 "Rows with script_status found / not_found are set by the script and need no verdict."]:
        ws2.append([line])
    wb.save(path)


def stop_checks(atlas_out: pd.DataFrame, epoch_out: pd.DataFrame, counts: pd.DataFrame, frame: P.Frame, expected_atlas: int, atlas_cols: list[str], epoch_cols: list[str]):
    """Section 2.9; every failure stops the build."""
    stops = []
    if len(atlas_out) != expected_atlas:
        stops.append(f"Atlas {len(atlas_out)} != {expected_atlas}")
    if len(epoch_out) != 94:
        stops.append(f"Epoch {len(epoch_out)} != 94")
    if len(counts) != 242:
        stops.append(f"ADM0 {len(counts)} != 242")
    if int(counts["n_records"].sum()) != expected_atlas:
        stops.append("sum of the country count != Atlas records")
    for nm, d in (("Atlas", atlas_out), ("Epoch", epoch_out)):
        if d["record_id"].duplicated().any():
            stops.append(f"{nm}: record_id not unique")
        unit_ids = set(frame.units["unit_id"])
        bad_u = d[(d["unit_id"] != "") & ~d["unit_id"].isin(unit_ids)]
        if len(bad_u):
            stops.append(f"{nm}: {len(bad_u)} unit_id not in the frame")
        isos = set(frame.units.loc[frame.units["level"] == "ADM0", "iso3"])
        for col in ("iso3_source", "iso3_frame", "control_iso3"):
            bad_i = d[(d[col] != "") & ~d[col].isin(isos)]
            if len(bad_i):
                stops.append(f"{nm}: {len(bad_i)} {col} not in the frame")
        city = d["coord_precision"] == "city"
        if not (d.loc[city, ["lon", "lat"]].astype(str).ne("").all().all() and (d.loc[city, "unit_id"] == "").all()):
            stops.append(f"{nm}: a city record without lon/lat or with unit_id")
        if (d.loc[city, ["place_name", "place_ref"]] == "").any().any() or not d.loc[city, "place_ref"].str.match(r"^(geonames:[0-9]+|census:[0-9]+)$").all():
            stops.append(f"{nm}: a city record without place_name / place_ref or with a malformed place_ref")
        non = ~city
        if (d.loc[non, ["place_name", "place_ref"]] != "").any().any():
            stops.append(f"{nm}: a county/none record with place_name / place_ref")
        if not ((d.loc[non, "unit_id"] != "").all() and (d.loc[non, "lon"].astype(str) == "").all() and d.loc[non, "coord_precision"].isin(["county", "none"]).all()):
            stops.append(f"{nm}: a county/none record without unit_id or with lon/lat")
        if (d["power_basis"] == "").any() or d["power_basis"].isna().any():
            stops.append(f"{nm}: a record without power_basis")
        if (d["ref"] == "").any():
            stops.append(f"{nm}: a record without ref (G10)")
        nl = d.astype(str).apply(lambda col: col.str.contains("[\r\n]", regex=True)).any()
        if nl.any():
            stops.append(f"{nm}: line breaks inside fields {list(nl[nl].index)} (replace with ' | ')")
        for col in ("location_state", "power_status", "power_status_reason", "power_basis", "coord_source", "coord_precision", "resolved_by_population"):
            if not d[col].astype(str).str.fullmatch(r"[a-z0-9_]*").all():
                stops.append(f"{nm}: {col} is not an English code (G7)")
    # G3 / G3b: one power basis per record; the numeric power_mw field in one layer only
    if "power_mw" in atlas_out.columns:
        stops.append("G3b: power_mw in the Atlas layer")
    if list(atlas_out.columns) != atlas_cols or list(epoch_out.columns) != epoch_cols:
        stops.append("columns differ from the pre-flight list")
    if stops:
        raise L.BuildStop("; ".join(stops))
